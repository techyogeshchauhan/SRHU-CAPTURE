"""The shared Excel download: values as cells want them, never a formula."""

import unittest
from datetime import date, datetime, timezone
from io import BytesIO

from openpyxl import load_workbook

from app.services.excel_export import build_workbook, event_rows, excel_value


class ExcelExportTests(unittest.TestCase):
    def test_text_that_looks_like_a_formula_stays_text(self):
        content = build_workbook("Events", ["Name", "Date", "Count"], [
            ["=HYPERLINK(\"http://evil\",\"x\")", "2026-03-10", 4],
        ])
        sheet = load_workbook(BytesIO(content)).active
        self.assertEqual(sheet["A1"].value, "Events")
        self.assertEqual([c.value for c in sheet[4]], ["Name", "Date", "Count"])
        cell = sheet["A5"]
        self.assertEqual(cell.data_type, "s")
        self.assertEqual(cell.value, '=HYPERLINK("http://evil","x")')
        self.assertEqual(sheet["B5"].value.date(), date(2026, 3, 10))
        self.assertEqual(sheet["C5"].value, 4)

    def test_values(self):
        self.assertEqual(excel_value(["Scopus", "SCI"]), "Scopus, SCI")
        self.assertIsNone(excel_value(""))
        aware = datetime(2026, 3, 10, 9, 30, tzinfo=timezone.utc)
        self.assertIsNone(excel_value(aware).tzinfo)
        self.assertEqual(excel_value("2026-03-10T09:30:00+00:00"), datetime(2026, 3, 10, 9, 30))

    def test_an_empty_list_says_so(self):
        sheet = load_workbook(BytesIO(build_workbook("Teachers", ["Name"], []))).active
        self.assertEqual(sheet["A5"].value, "Nothing matches these filters.")

    def test_event_rows_take_the_department_from_the_form_metadata(self):
        rows = event_rows([{
            "event_name": "Talk", "status": "in_progress", "teacher_name": "Asha",
            "description": 'Talk.\n\n<!--CC_METADATA:{"department":"SST"}-->',
        }])
        self.assertEqual(rows[0][0], "Talk")
        self.assertEqual(rows[0][2], "In progress")
        self.assertEqual(rows[0][11], "SST")
        self.assertEqual(rows[0][-1], "Talk.")


if __name__ == "__main__":
    unittest.main()
