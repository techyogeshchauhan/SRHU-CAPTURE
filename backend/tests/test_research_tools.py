"""Research tools: analytics, exports, CVs, import, DOI, reminders, duplicates."""

from __future__ import annotations

import copy
import unittest
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from unittest.mock import patch

import httpx
from openpyxl import load_workbook

from app.services.research_doi import DoiError, lookup_doi, parse_crossref
from app.services.research_identifiers import identifiers_of, normalize_doi
from app.services.research_reminders import reminders_for
from app.services.research_spreadsheets import import_template
from tests.test_research import MANAGER_ID, OTHER_TEACHER_ID, PATENT, TEACHER_ID, ResearchFixture

JOURNAL = {
    "category": "journal",
    "title": "Herbal dressings for chronic wounds in rural clinics",
    "record_date": "2026-03-10",
    "outcome_status": "published",
    "details": {
        "authors": "Meera Joshi, Arjun Rawat",
        "author_position": "First author",
        "journal": "Wound Care",
        "indexing": ["Scopus", "PubMed"],
        "doi": "https://doi.org/10.1016/j.wc.2026.01.001",
    },
}

CROSSREF = {
    "type": "journal-article",
    "title": ["Herbal <i>dressings</i> for chronic wounds"],
    "container-title": ["Wound Care"],
    "publisher": "Elsevier",
    "volume": "12",
    "issue": "3",
    "page": "45-60",
    "ISSN": ["1234-5678"],
    "author": [{"given": "Meera", "family": "Joshi"}, {"given": "Arjun", "family": "Rawat"}],
    "published-print": {"date-parts": [[2026, 3]]},
    "URL": "https://doi.org/10.1016/j.wc.2026.01.001",
}


class ToolsTest(ResearchFixture, unittest.TestCase):
    def verified(self, who="manager", **overrides):
        """An accepted record: the Event Manager's is recorded on submission."""
        record = self.create(who, **overrides)
        self.attach(record["id"], who, tag=self.tag_for(overrides.get("category", "patent")))
        done = self.call("POST", f"/research/records/{record['id']}/submit", who)["record"]
        if done["verification_status"] != "verified":
            done = self.call("PATCH", f"/dean/research/records/{record['id']}/verify", "dean")["record"]
        return done

    @staticmethod
    def tag_for(category):
        return {"patent": "Filing receipt", "journal": "Full paper"}[category]

    def download(self, path, who, expect=200):
        response = self.client.get(path, headers={"Authorization": f"Bearer {who}"})
        self.assertEqual(response.status_code, expect, response.text[:300])
        return response

    # ------------------------------------------------------------ analytics

    def test_analytics_count_verified_unless_asked(self):
        self.verified()
        self.verified(**copy.deepcopy(JOURNAL))
        self.submitted()
        verified = self.call("GET", "/dean/research/analytics", "dean")
        self.assertEqual(verified["total"], 2)
        self.assertEqual(verified["by_category"]["patent"], 1)
        self.assertEqual(verified["indexing"], {"Scopus": 1, "PubMed": 1})
        self.assertEqual(verified["top_contributors"][0]["name"], "Meera Joshi")
        # One school (SST): no department breakdown to show.
        self.assertNotIn("by_department", verified)
        self.assertNotIn("department", verified["top_contributors"][0])
        pending = self.call("GET", "/dean/research/analytics?include_pending=true", "dean")
        self.assertEqual(pending["total"], 3)
        self.assertEqual(self.call("GET", "/research/analytics", "manager")["total"], 2)
        self.call("GET", "/research/analytics", "teacher", 403)

    # ------------------------------------------------------------ exports

    def test_register_export_respects_who_is_asking(self):
        self.verified()
        self.create(faculty_id=str(OTHER_TEACHER_ID), title="Another teacher's draft patent filing")
        book = load_workbook(BytesIO(self.download("/research/export.xlsx", "teacher").content))
        self.assertEqual(book["Patent"].max_row, 2)  # header + the teacher's one record
        book = load_workbook(BytesIO(self.download("/research/export.xlsx", "manager").content))
        self.assertEqual(book["Patent"].max_row, 3)
        book = load_workbook(BytesIO(self.download("/dean/research/export.xlsx", "dean").content))
        self.assertEqual(book["Patent"].max_row, 2)  # the Dean never sees drafts
        headers = [cell.value for cell in book["Patent"][1]]
        self.assertIn("Application number", headers)
        self.assertNotIn("Department", headers)

    def test_naac_export_holds_verified_records_only(self):
        self.verified(**copy.deepcopy(JOURNAL))
        self.submitted()
        response = self.download("/dean/research/export/naac.xlsx?academic_year=2025-26", "dean")
        self.assertIn("naac-criterion-3-research-2025-26.xlsx", response.headers["content-disposition"])
        book = load_workbook(BytesIO(response.content))
        self.assertIn("Research papers", book.sheetnames)
        papers = book["Research papers"]
        self.assertEqual(papers.cell(row=3, column=2).value, JOURNAL["title"])
        # NAAC asks for the teacher's department: always the one school.
        self.assertEqual(papers.cell(row=3, column=4).value, "School of Science and Technology")
        self.assertEqual(papers.cell(row=3, column=10).value, "https://doi.org/10.1016/j.wc.2026.01.001")
        patents = book["Patents"]
        self.assertEqual(patents.cell(row=3, column=2).value, "No verified records of this kind for the period.")
        self.download("/research/export/naac.xlsx", "teacher", 403)

    # ------------------------------------------------------------ CV

    def test_research_cv(self):
        self.verified()
        pdf = self.download(f"/research/faculty/{TEACHER_ID}/cv.pdf", "teacher")
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        self.download(f"/research/faculty/{OTHER_TEACHER_ID}/cv.pdf", "teacher", 403)
        self.download(f"/research/faculty/{OTHER_TEACHER_ID}/cv.pdf", "manager", 403)  # not the Event Manager's
        self.download(f"/dean/research/faculty/{TEACHER_ID}/cv.pdf", "dean")
        self.download(f"/dean/research/faculty/{MANAGER_ID}/cv.pdf", "dean", 404)

    # ------------------------------------------------------------ import

    def filled_template(self):
        book = load_workbook(BytesIO(import_template()))
        journal = book["Journal"]
        headers = [cell.value for cell in journal[1]]

        def row(sheet, values):
            cells = [None] * len([c for c in sheet[1]])
            for header, value in values.items():
                cells[[c.value for c in sheet[1]].index(header)] = value
            sheet.append(cells)

        row(journal, {
            "Teacher email *": "MEERA@srhu.edu.in",
            "Paper title *": "Herbal dressings for chronic wounds in rural clinics",
            "Publication date * (YYYY-MM-DD)": datetime(2026, 3, 10),
            "Status": "Published",
            "Faculty member's authorship *": "First author",
            "Journal *": "Wound Care",
            "Indexed in * (separate with ;)": "Scopus; PubMed",
        })
        row(journal, {
            "Teacher email *": "nobody@srhu.edu.in",
            "Paper title *": "A paper by someone without an account",
            "Publication date * (YYYY-MM-DD)": "10/03/2026",
        })
        row(journal, {
            "Teacher email *": "arjun@srhu.edu.in",
            "Paper title *": "Quartile typo paper about wound care outcomes",
            "Publication date * (YYYY-MM-DD)": "2026-02-01",
            "Quartile": "Q9",
        })
        patent = book["Patent"]
        row(patent, {
            "Teacher email *": "arjun@srhu.edu.in",
            "Title of the invention *": "Solar sterilising cabinet for clinics",
            "Filing date * (YYYY-MM-DD)": "2024-01-15",
            "Status": "filed",
            "Inventors *": "Arjun Rawat",
            "Patent office *": "India (IPO)",
        })
        self.assertIn("Teacher email *", headers)
        out = BytesIO()
        book.save(out)
        return out.getvalue()

    def test_import_previews_then_creates_drafts(self):
        self.download("/research/import/template.xlsx", "manager")
        self.download("/research/import/template.xlsx", "teacher", 403)
        preview = self.call(
            "POST", "/research/import/preview", "manager",
            files={"file": ("research.xlsx", self.filled_template(), "application/octet-stream")},
        )
        self.assertEqual((preview["total"], preview["valid"], preview["invalid"]), (4, 2, 2))
        by_title = {row["title"]: row for row in preview["rows"]}
        self.assertEqual(by_title["Herbal dressings for chronic wounds in rural clinics"]["faculty"]["name"], "Meera Joshi")
        self.assertIn("No teacher account", by_title["A paper by someone without an account"]["errors"][0])
        self.assertIn("Q9", by_title["Quartile typo paper about wound care outcomes"]["errors"][0])
        patent = by_title["Solar sterilising cabinet for clinics"]
        self.assertIn("Application number", patent["warnings"][0])

        rows = [row["payload"] for row in preview["rows"] if row["payload"]]
        # The server checks again: a row tampered with after the preview fails.
        rows.append({**rows[0], "faculty_id": str(MANAGER_ID), "title": "Smuggled onto a non-teacher account"})
        result = self.call("POST", "/research/import", "manager", json={"rows": rows})
        self.assertEqual(len(result["created"]), 2)
        self.assertEqual(len(result["failed"]), 1)
        self.assertTrue(all(doc["verification_status"] == "draft" for doc in self.records.docs))
        self.assertEqual(self.records.docs[0]["history"][0]["note"], "Imported from a spreadsheet")
        # One notification per teacher, however many rows were theirs.
        self.assertEqual(len(self.notes_for(TEACHER_ID, "research_added")), 1)
        self.assertEqual(len(self.notes_for(OTHER_TEACHER_ID, "research_added")), 1)

    def test_import_refuses_what_is_not_a_template(self):
        self.call("POST", "/research/import/preview", "manager", 400,
                  files={"file": ("notes.xlsx", b"not a workbook", "application/octet-stream")})

    # ------------------------------------------------------------ DOI

    def test_crossref_answer_becomes_a_journal_record(self):
        found = parse_crossref(CROSSREF, "10.1016/j.wc.2026.01.001")
        self.assertEqual(found["category"], "journal")
        self.assertEqual(found["title"], "Herbal dressings for chronic wounds")
        self.assertEqual(found["record_date"], "2026-03-01")
        self.assertEqual(found["details"]["authors"], "Meera Joshi, Arjun Rawat")
        self.assertEqual(found["details"]["pages"], "45-60")

    # Trimmed from real Crossref answers, each checked against the paper.

    def test_journal_date_is_when_the_paper_was_published(self):
        # NumPy, Nature 585: "Published 16 September 2020", in the 17 September issue.
        found = parse_crossref({
            **CROSSREF,
            "published-print": {"date-parts": [[2020, 9, 17]]},
            "published-online": {"date-parts": [[2020, 9, 16]]},
            "issued": {"date-parts": [[2020, 9, 16]]},
            "created": {"date-parts": [[2020, 9, 10]]},
            "ISSN": ["1476-4687", "0028-0836"],
            "issn-type": [{"value": "1476-4687", "type": "electronic"}, {"value": "0028-0836", "type": "print"}],
        }, "10.1038/s41586-020-2649-2")
        self.assertEqual(found["record_date"], "2020-09-16")
        self.assertEqual(found["notes"], [])
        self.assertEqual(found["details"]["issn"], "0028-0836")

    def test_a_rough_date_is_flagged_not_passed_off_as_exact(self):
        found = parse_crossref(CROSSREF, "10.1016/j.wc.2026.01.001")
        self.assertEqual(found["record_date"], "2026-03-01")
        self.assertIn("only the month of publication (March 2026)", found["notes"][0])
        # A year-only date that is more precise online is taken from there.
        found = parse_crossref({**CROSSREF, "published-print": {"date-parts": [[2015]]},
                                "published-online": {"date-parts": [[2015, 11, 18]]}}, "10.1/x")
        self.assertEqual((found["record_date"], found["notes"]), ("2015-11-18", []))

    def test_title_keeps_its_subtitle_and_plain_text(self):
        # XGBoost, KDD '16: Crossref splits "XGBoost" / "A Scalable Tree Boosting System".
        found = parse_crossref({**CROSSREF, "title": ["XGBoost"], "subtitle": ["A Scalable Tree Boosting System"]}, "10.1/x")
        self.assertEqual(found["title"], "XGBoost: A Scalable Tree Boosting System")
        found = parse_crossref({**CROSSREF, "title": ["Fitting Linear Mixed-Effects Models Using\n   <b>lme4</b>"],
                                "subtitle": []}, "10.1/x")
        self.assertEqual(found["title"], "Fitting Linear Mixed-Effects Models Using lme4")
        found = parse_crossref({**CROSSREF, "title": ["Wounds &amp; dressings"],
                                "author": [{"given": "B. P.", "family": "Abbott"}]}, "10.1/x")
        self.assertEqual(found["title"], "Wounds & dressings")
        self.assertEqual(found["details"]["authors"], "B. P. Abbott")

    def test_article_number_stands_in_for_pages(self):
        # Scientific Reports 7, 46130 has no page range.
        found = parse_crossref({**CROSSREF, "page": None, "article-number": "46130"}, "10.1038/srep46130")
        self.assertEqual(found["details"]["pages"], "Article 46130")

    def test_abstract_is_plain_paragraphs(self):
        found = parse_crossref({**CROSSREF, "abstract": (
            "<jats:title>Abstract</jats:title><jats:p>Used in the discovery of gravitational waves\n"
            "<jats:sup>1</jats:sup>\n. Sequenced the (<jats:italic>16S rRNA</jats:italic>) gene of H<jats:sub>2</jats:sub>O "
            "at a strain of <mml:math><mml:mn>1.0</mml:mn><mml:mo>×</mml:mo><mml:msup><mml:mrow><mml:mn>10</mml:mn></mml:mrow>"
            "<mml:mrow><mml:mo>−</mml:mo><mml:mn>21</mml:mn></mml:mrow></mml:msup></mml:math>.</jats:p>"
            "<jats:sec><jats:title>Results</jats:title><jats:p>It &amp; worked.</jats:p></jats:sec>")}, "10.1/x")
        self.assertEqual(found["abstract"], "Used in the discovery of gravitational waves. Sequenced the (16S rRNA) "
                         "gene of H2O at a strain of 1.0×10^−21.\n\nResults: It & worked.")
        long = parse_crossref({**CROSSREF, "abstract": "<jats:p>" + "word " * 2000 + "</jats:p>"}, "10.1/x")["abstract"]
        self.assertLessEqual(len(long), 4000)
        self.assertTrue(long.endswith("word…"))

    def test_a_thousand_authors_are_shortened_to_fit(self):
        people = [{"given": f"A{n}", "family": f"Author{n}"} for n in range(1012)]
        found = parse_crossref({**CROSSREF, "author": people}, "10.1103/PhysRevLett.116.061102")
        self.assertLessEqual(len(found["details"]["authors"]), 4000)
        self.assertTrue(found["details"]["authors"].endswith("et al. (1012 authors in all)"))
        self.assertEqual(len(found["author_names"]), 1012)
        self.assertTrue(any("1012 authors" in note for note in found["notes"]))

    def test_conference_paper_takes_the_conference_dates(self):
        # ResNet, CVPR 2016: Las Vegas, 27-30 June 2016.
        found = parse_crossref({
            "type": "proceedings-article",
            "title": ["Deep Residual Learning for Image Recognition"],
            "container-title": ["2016 IEEE Conference on Computer Vision and Pattern Recognition (CVPR)"],
            "event": {"name": "2016 IEEE Conference on Computer Vision and Pattern Recognition (CVPR)",
                      "location": "Las Vegas, NV, USA", "start": {"date-parts": [[2016, 6, 27]]},
                      "end": {"date-parts": [[2016, 6, 30]]}},
            "page": "770-778", "published-print": {"date-parts": [[2016, 6]]},
            "author": [{"given": "Kaiming", "family": "He"}],
        }, "10.1109/cvpr.2016.90")
        self.assertEqual(found["category"], "conference")
        self.assertEqual(found["record_date"], "2016-06-27")
        self.assertEqual(found["details"]["end_date"], "2016-06-30")
        self.assertEqual(found["details"]["venue"], "Las Vegas, NV, USA")
        self.assertEqual(found["notes"], [])
        self.assertIn("organiser", found["doi_fields"])  # cleared, though Crossref has none

        # KDD '16 gives no dates: the proceedings' date stands in, with a note.
        found = parse_crossref({
            "type": "proceedings-article", "title": ["XGBoost"],
            "event": {"name": "KDD '16", "sponsor": ["SIGMOD ACM", "SIGKDD ACM"]},
            "published-print": {"date-parts": [[2016, 8, 13]]},
        }, "10.1145/2939672.2939785")
        self.assertEqual(found["record_date"], "2016-08-13")
        self.assertIn("does not give the conference dates", found["notes"][0])
        self.assertEqual(found["details"]["organiser"], "SIGMOD ACM; SIGKDD ACM")

    def test_chapter_names_its_book_not_the_series(self):
        # U-Net: a chapter of MICCAI 2015, in Lecture Notes in Computer Science.
        found = parse_crossref({
            "type": "book-chapter",
            "title": ["U-Net: Convolutional Networks for Biomedical Image Segmentation"],
            "container-title": ["Lecture Notes in Computer Science",
                                "Medical Image Computing and Computer-Assisted Intervention – MICCAI 2015"],
            "publisher": "Springer International Publishing", "page": "234-241",
            "ISBN": ["9783319245744", "9783319245737"],
            "isbn-type": [{"value": "9783319245744", "type": "electronic"}, {"value": "9783319245737", "type": "print"}],
            "published-print": {"date-parts": [[2015]]}, "published-online": {"date-parts": [[2015, 11, 18]]},
        }, "10.1007/978-3-319-24574-4_28")
        self.assertEqual(found["details"]["book_title"], "Medical Image Computing and Computer-Assisted Intervention – MICCAI 2015")
        self.assertEqual(found["details"]["isbn"], "9783319245737")
        self.assertEqual(found["record_date"], "2015-11-18")
        self.assertTrue(any("conference proceedings" in note for note in found["notes"]))

        book = parse_crossref({"type": "book", "title": ["Statistical Rethinking"], "edition-number": "2",
                               "container-title": ["Springer Series in Statistics"], "publisher": "CRC"}, "10.1/x")
        self.assertEqual(book["details"]["edition"], "2nd")
        self.assertNotIn("book_title", book["details"])

    def test_preprint_is_not_passed_off_as_published(self):
        found = parse_crossref({
            "type": "posted-content", "title": ["A SARS-CoV-2-Human Protein-Protein Interaction Map"],
            "container-title": [], "institution": [{"name": "bioRxiv"}],
            "posted": {"date-parts": [[2020, 3, 22]]},
            "relation": {"is-preprint-of": [{"id": "10.1038/s41586-020-2286-9", "id-type": "doi"}]},
        }, "10.1101/2020.03.22.002386")
        self.assertIsNone(found["outcome_status"])
        self.assertEqual(found["details"]["journal"], "bioRxiv")
        self.assertIn("10.1038/s41586-020-2286-9", found["notes"][0])

    def test_doi_lookup_reports_crossref_failures_plainly(self):
        def answer(status_code, body=None):
            transport = httpx.MockTransport(lambda request: httpx.Response(status_code, json=body or {}))
            return httpx.Client(transport=transport)

        with self.assertRaises(DoiError) as missing:
            lookup_doi("10.9999/not.there", client=answer(404))
        self.assertEqual(missing.exception.status, 404)
        with self.assertRaises(DoiError):
            lookup_doi("not a doi")
        found = lookup_doi("doi:10.1016/J.WC.2026.01.002", client=answer(200, {"message": CROSSREF}))
        self.assertEqual(found["doi"], "10.1016/j.wc.2026.01.002")

    def test_doi_endpoint_fits_the_suggestion_to_the_form(self):
        suggestion = parse_crossref(CROSSREF, "10.1016/j.wc.2026.01.001")
        suggestion["details"]["quartile"] = "Q9"  # not a choice the form has
        with patch("app.routers.research_tools.lookup_doi", return_value=suggestion):
            body = self.call("GET", "/research/doi?doi=10.1016/j.wc.2026.01.001", "teacher")["suggestion"]
        self.assertNotIn("quartile", body["details"])
        self.assertEqual(body["details"]["journal"], "Wound Care")

    # ------------------------------------------------------------ duplicates

    def test_a_doi_is_found_in_whatever_form_it_is_pasted(self):
        """"DOI: https://doi.org/…" was refused as "does not look like a DOI"."""
        bare = "10.1103/physrevlett.116.061102"
        for typed in (
            "DOI: https://doi.org/10.1103/PhysRevLett.116.061102",
            "https://doi.org/10.1103/PhysRevLett.116.061102",
            "https://dx.doi.org/10.1103/PhysRevLett.116.061102?utm_source=x#top",
            "doi:10.1103/PhysRevLett.116.061102",
            "DOI 10.1103/PhysRevLett.116.061102",
            "  10.1103/PhysRevLett.116.061102.  ",
            "https://journals.aps.org/prl/abstract/10.1103/PhysRevLett.116.061102",
        ):
            self.assertEqual(normalize_doi(typed), bare, typed)
        # Brackets that belong to the DOI stay; one closing the sentence goes.
        self.assertEqual(normalize_doi("https://doi.org/10.1016/S0140-6736(20)30183-5"), "10.1016/s0140-6736(20)30183-5")
        self.assertEqual(normalize_doi("(see 10.1038/nature14539)."), "10.1038/nature14539")
        self.assertEqual(normalize_doi("https://doi.org/10.1002%2Fanie.201915678"), "10.1002/anie.201915678")
        # A Nature article page names the DOI's suffix only.
        self.assertEqual(normalize_doi("https://www.nature.com/articles/nature14236"), "10.1038/nature14236")
        self.assertEqual(normalize_doi("Publisher page: https://www.nature.com/articles/s41586-020-2649-2."),
                         "10.1038/s41586-020-2649-2")
        for junk in ("not a doi", "https://example.com/paper", "", "https://doi.org/", "10.", "10.1038",
                     "https://doi.org/invalid", "DOI: not available"):
            self.assertIsNone(normalize_doi(junk), junk)
        transport = httpx.MockTransport(lambda request: httpx.Response(200, json={"message": CROSSREF}))
        found = lookup_doi("DOI: https://doi.org/10.1016/J.WC.2026.01.003", client=httpx.Client(transport=transport))
        self.assertEqual(found["doi"], "10.1016/j.wc.2026.01.003")

    def test_identifiers_normalise(self):
        self.assertEqual(normalize_doi("https://doi.org/10.1016/ABC"), "10.1016/abc")
        keys = identifiers_of("book", "A chapter on wound care practice", {"book_type": "Book chapter", "isbn": "978-0-12-345678-9"})
        self.assertEqual(keys[0], "chapter:9780123456789:a chapter on wound care practice")

    def test_the_same_paper_twice_is_flagged_not_refused(self):
        first = self.create(**copy.deepcopy(JOURNAL))
        second = self.create(faculty_id=str(OTHER_TEACHER_ID), **{**copy.deepcopy(JOURNAL), "title": "Different words entirely here"})
        again = self.call("GET", f"/research/records/{second['id']}", "manager")["record"]
        self.assertEqual(again["duplicates"][0]["id"], first["id"])
        self.assertEqual(again["duplicates"][0]["match"], "DOI")
        self.assertEqual(again["duplicates"][0]["faculty_name"], "Meera Joshi")
        self.assertFalse(again["duplicates"][0]["same_faculty"])

        live = self.call("POST", "/research/duplicates", "teacher", json={
            "category": "journal", "title": "x", "details": {"doi": "10.1016/j.wc.2026.01.001"},
        })["duplicates"]
        # The teacher is warned about their own draft only; Arjun's draft is his business.
        self.assertEqual([d["id"] for d in live], [first["id"]])

    # ------------------------------------------------------------ reminders

    def test_outcomes_due_and_reminders_are_not_repeated(self):
        record = self.create(record_date=(date.today() - timedelta(days=600)).isoformat())
        due = self.call("GET", "/research/outcomes-due", "teacher")
        self.assertEqual(due["items"][0]["id"], record["id"])
        self.assertEqual(due["items"][0]["reminders"][0]["key"], "patent-publication")

        self.assertEqual(self.call("POST", "/research/reminders/check", "teacher")["created"], 1)
        self.assertEqual(self.call("POST", "/research/reminders/check", "teacher")["created"], 0)
        # The Event Manager gets a weekly digest, not one per record.
        self.assertEqual(self.call("POST", "/research/reminders/check", "manager")["created"], 1)
        self.assertEqual(self.call("POST", "/research/reminders/check", "manager")["created"], 0)

    def test_reminder_rules(self):
        today = date(2026, 10, 10)
        project = {"category": "project", "outcome_status": "ongoing", "details": {"end_date": "2026-09-30"},
                   "verification_status": "verified"}
        self.assertEqual(reminders_for(project, today), [])
        draft = {"category": "other", "verification_status": "draft",
                 "updated_at": datetime(2026, 8, 1, tzinfo=timezone.utc)}
        self.assertEqual([r.key for r in reminders_for(draft, today)], ["draft-idle"])
        self.assertEqual(reminders_for({**project, "verification_status": "rejected"}, today), [])

    # ------------------------------------------------------------ limits

    def test_evidence_limits_come_from_the_super_admin_settings(self):
        limits = {"max_research_evidence_files": 1, "max_research_evidence_total_mb": 10,
                  "max_research_evidence_total_bytes": 10 * 1024 * 1024}
        with patch("app.routers.research.get_upload_limits", lambda: limits):
            record = self.create()
            self.attach(record["id"])
            self.attach(record["id"], expect=400)
            self.assertEqual(self.call("GET", "/research/categories", "teacher")["limits"]["max_evidence_files"], 1)

    def test_super_admin_schema_accepts_and_bounds_the_evidence_limits(self):
        from pydantic import ValidationError

        from app.schemas.superadmin import UploadLimitsUpdateRequest

        base = {
            "max_photos_per_event": 10, "max_photo_size_mb": 20, "max_video_size_mb": 200,
            "max_video_total_mb": 200, "max_documents_per_event": 50, "max_documents_total_mb": 15,
        }
        defaults = UploadLimitsUpdateRequest(**base)
        self.assertEqual((defaults.max_research_evidence_files, defaults.max_research_evidence_total_mb), (5, 10))
        with self.assertRaises(ValidationError):
            UploadLimitsUpdateRequest(**base, max_research_evidence_files=0)


class AddFacultyTest(ResearchFixture, unittest.TestCase):
    """The Event Manager adds a faculty member from the research form."""

    def setUp(self):
        super().setUp()
        # The lookup and the invitation write go through the auth and Dean
        # routers' own handles; point them at the same fake.
        for target in ("app.routers.auth.users", "app.routers.dean_teachers.users"):
            patcher = patch(target, self.users)
            patcher.start()
            self.addCleanup(patcher.stop)

    def add(self, who="manager", expect=201, **fields):
        body = {"name": "Dr. Kavita Negi", "email": "kavita.negi@srhu.edu.in",
                "designation": "Assistant Professor", "send_invite": False, **fields}
        return self.call("POST", "/research/faculty", who, expect, json=body)

    def test_adds_a_teacher_account_the_dean_would_recognise(self):
        body = self.add()
        self.assertFalse(body["existing"])
        doc = next(d for d in self.users.docs if d["email"] == "kavita.negi@srhu.edu.in")
        self.assertEqual(doc["role"], "teacher")
        self.assertTrue(doc["must_change_password"])  # no usable password until invited
        self.assertFalse(doc["email_verified"])
        self.assertEqual((doc["onboarded_via"], doc["added_by_role"]), ("dean_added", "event_manager"))
        self.assertEqual(doc["department"], "SST")  # the one school the platform serves
        self.assertNotIn("department", body["faculty"])
        # ... and research can be recorded for them straight away.
        record = self.create(faculty_id=body["faculty"]["id"])
        self.assertEqual(record["faculty"]["name"], "Dr. Kavita Negi")

    def test_an_existing_teacher_is_selected_not_duplicated(self):
        body = self.add(expect=200, email="MEERA@srhu.edu.in")
        self.assertTrue(body["existing"])
        self.assertEqual(body["faculty"]["id"], str(TEACHER_ID))
        self.assertEqual(sum(1 for d in self.users.docs if d["email"] == "meera@srhu.edu.in"), 1)

    def test_other_accounts_are_refused_and_only_the_manager_may_add(self):
        self.add(expect=409, email="dean@srhu.edu.in")
        self.add("teacher", expect=403)
        self.add("dean", expect=403)

    def test_the_invitation_says_who_added_them(self):
        sent = []
        with patch("app.routers.research.email_service.is_configured", lambda: True), patch(
            "app.services.email_service.send_teacher_invitation_email",
            lambda to, name, token, **wording: sent.append((to, wording)),
        ):
            body = self.add(send_invite=True)
        self.assertEqual(body["invite"]["status"], "sent")
        self.assertEqual(sent[0][0], "kavita.negi@srhu.edu.in")
        self.assertEqual(sent[0][1]["added_by"], "The Event Manager")


PATENT_DETAILS = PATENT["details"]

if __name__ == "__main__":
    unittest.main()
