"""Research records to and from Excel workbooks.

* ``export_records``   -- the filtered register, one sheet per category, every
                          field in its own column.
* ``export_naac``      -- verified records laid out the way NAAC Criterion III
                          (Research, Innovations and Extension) data templates
                          ask for them, one sheet per kind of output.
* ``import_template``  -- the workbook an Event Manager fills in to bring
                          existing records in, one sheet per category.
* ``parse_import``     -- that workbook back into plain rows. Validation is the
                          router's job (it needs the teachers' accounts).

openpyxl only, already a dependency. Dates go out as real date cells so
Excel sorts and filters them as dates.
"""

from __future__ import annotations

import re
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.config import settings
from app.services.research_categories import CATEGORIES, Category, indexing_text, outcome_label

# Every researcher belongs to the one school the platform serves. NAAC's
# templates still ask for the teacher's department, so it is filled from here.
SCHOOL = settings.report_school

_BRAND = "0D3668"
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor=_BRAND)
_TITLE_FONT = Font(bold=True, size=13, color=_BRAND)
_NOTE_FONT = Font(italic=True, color="5B6475")
_WRAP = Alignment(wrap_text=True, vertical="top")

VERIFICATION_LABEL = {
    "draft": "Draft",
    "submitted": "Awaiting verification",
    "verified": "Verified",
    "changes_requested": "Changes requested",
    "rejected": "Rejected",
}

EMAIL_HEADER = "Teacher email"
MAX_IMPORT_ROWS = 500


# ------------------------------------------------------------ helpers

def _to_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        try:
            return date.fromisoformat(text)
        except ValueError:
            return None
    return None


def _cell(value, field_type: str | None = None):
    if value is None or value == "":
        return None
    if isinstance(value, list):
        return ", ".join(str(v) for v in value)
    if field_type == "date":
        return _to_date(value) or value
    return value


def _finish(sheet, header_row: int, widths: dict[int, int] | None = None) -> None:
    """Style the header row, size the columns, freeze below the header."""
    for cell in sheet[header_row]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    for index, column in enumerate(sheet.iter_cols(min_row=header_row), start=1):
        longest = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        sheet.column_dimensions[get_column_letter(index)].width = (widths or {}).get(index) or min(max(10, longest + 2), 60)
        for c in column[1:]:
            if isinstance(c.value, (date, datetime)):
                c.number_format = "dd-mmm-yyyy"
            elif isinstance(c.value, str) and len(c.value) > 40:
                c.alignment = _WRAP


def _sheet_name(text: str) -> str:
    return re.sub(r"[\[\]:*?/\\]", "-", text)[:31]


def _person(people: dict, faculty_id) -> dict:
    return people.get(str(faculty_id or "")) or {}


def _save(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _year(iso: str | None) -> int | None:
    found = _to_date(iso)
    return found.year if found else None


def _iso(stamp) -> date | None:
    if isinstance(stamp, datetime):
        return stamp.date()
    return _to_date(str(stamp or "")[:10])


# ------------------------------------------------------------ register export

def export_records(records: list[dict], people: dict, *, heading: str) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["A1"] = heading
    summary["A1"].font = _TITLE_FONT
    summary["A2"] = f"Exported {datetime.now().strftime('%d %b %Y, %H:%M')} · {len(records)} records"
    summary["A2"].font = _NOTE_FONT
    summary.append([])
    summary.append(["Category", *VERIFICATION_LABEL.values(), "Total"])
    for category in CATEGORIES.values():
        rows = [r for r in records if r.get("category") == category.key]
        summary.append([
            category.label,
            *[sum(1 for r in rows if r.get("verification_status") == key) for key in VERIFICATION_LABEL],
            len(rows),
        ])
    _finish(summary, 4)

    for category in CATEGORIES.values():
        rows = [r for r in records if r.get("category") == category.key]
        if not rows:
            continue
        sheet = workbook.create_sheet(_sheet_name(category.label))
        sheet.append([
            "Faculty member", category.title_label, category.date_label, "Academic year",
            "Status", "Verification",
            *[field.label for field in category.fields],
            "Link", "Abstract", "Evidence files", "Entered by", "Verified on",
        ])
        for record in rows:
            person = _person(people, record.get("faculty_id"))
            details = record.get("details") or {}
            sheet.append([
                person.get("name"),
                record.get("title"),
                _to_date(record.get("record_date")),
                record.get("academic_year"),
                outcome_label(category, record.get("outcome_status")),
                VERIFICATION_LABEL.get(record.get("verification_status"), record.get("verification_status")),
                *[_cell(details.get(field.key), field.type) for field in category.fields],
                record.get("link"),
                record.get("abstract"),
                len(record.get("evidence") or []),
                record.get("created_by_name"),
                _iso(record.get("verified_at")) if record.get("verification_status") == "verified" else None,
            ])
        _finish(sheet, 1)
    return _save(workbook)


# ------------------------------------------------------------ NAAC export

def _naac_sheet(workbook: Workbook, name: str, caption: str, headers: list[str], rows: list[list]) -> None:
    sheet = workbook.create_sheet(_sheet_name(name))
    sheet["A1"] = caption
    sheet["A1"].font = _TITLE_FONT
    sheet.append(["Sl. No.", *headers])
    for number, row in enumerate(rows, start=1):
        sheet.append([number, *row])
    if not rows:
        sheet.append(["", "No verified records of this kind for the period."])
    _finish(sheet, 2)


def export_naac(records: list[dict], people: dict, *, period: str) -> bytes:
    """Verified records only -- the caller filters -- in NAAC's shape."""
    workbook = Workbook()
    cover = workbook.active
    cover.title = "About"
    cover["A1"] = "Research output for NAAC Criterion III (Research, Innovations and Extension)"
    cover["A1"].font = _TITLE_FONT
    cover["A2"] = f"Period: {period} · Verified records only · Exported {datetime.now().strftime('%d %b %Y')}"
    cover["A2"].font = _NOTE_FONT
    cover["A4"] = (
        "Each sheet lists one kind of output with the columns NAAC's data templates ask for. "
        "Every row was verified by the Dean against its evidence in Campus Capture."
    )
    cover["A4"].alignment = _WRAP
    cover.column_dimensions["A"].width = 110

    def by(category: str) -> list[dict]:
        found = [r for r in records if r.get("category") == category]
        return sorted(found, key=lambda r: r.get("record_date") or "")

    def teacher(record: dict) -> str:
        return _person(people, record.get("faculty_id")).get("name") or ""

    def d(record: dict, key: str):
        return (record.get("details") or {}).get(key)

    journals = by("journal")
    _naac_sheet(
        workbook, "Research papers", f"Research papers published in journals — {period}",
        ["Title of paper", "Name of the author/s", "Department of the teacher", "Name of journal",
         "Year of publication", "ISSN number", "Indexed in", "Impact factor", "DOI / link"],
        [[
            r.get("title"), d(r, "authors") or teacher(r), SCHOOL, d(r, "journal"),
            _year(r.get("record_date")), d(r, "issn"), indexing_text(r.get("details") or {}), d(r, "impact_factor"),
            f"https://doi.org/{d(r, 'doi')}" if d(r, "doi") else r.get("link"),
        ] for r in journals],
    )

    books = by("book")
    _naac_sheet(
        workbook, "Books and chapters", f"Books and chapters in edited volumes — {period}",
        ["Name of the teacher", "Type", "Title of the book / chapter", "Title of the book (for a chapter)",
         "National / International", "Year of publication", "ISBN number", "Name of the publisher"],
        [[
            teacher(r), d(r, "book_type"), r.get("title"), d(r, "book_title"), d(r, "level"),
            _year(r.get("record_date")), d(r, "isbn"), d(r, "publisher"),
        ] for r in books],
    )

    conferences = by("conference")
    _naac_sheet(
        workbook, "Conference papers", f"Papers in conferences and their proceedings — {period}",
        ["Name of the teacher", "Title of the paper", "Name of the conference", "Organised by",
         "National / International", "Type", "Date", "Title of the proceedings", "ISBN / ISSN", "DOI"],
        [[
            teacher(r), r.get("title"), d(r, "conference"), d(r, "organiser"), d(r, "level"),
            d(r, "presentation_type"), _to_date(r.get("record_date")), d(r, "proceedings"), d(r, "isbn"), d(r, "doi"),
        ] for r in conferences],
    )

    patents = by("patent")
    patent_category = CATEGORIES["patent"]
    _naac_sheet(
        workbook, "Patents", f"Patents filed, published and awarded — {period}",
        ["Name of the teacher", "Title of the patent", "Patent office", "Application number", "Status",
         "Date of filing", "Date of publication", "Date of award", "Patent number"],
        [[
            teacher(r), r.get("title"), d(r, "office"), d(r, "application_number"),
            outcome_label(patent_category, r.get("outcome_status")), _to_date(r.get("record_date")),
            _to_date(d(r, "publication_date")), _to_date(d(r, "grant_date")), d(r, "patent_number"),
        ] for r in patents],
    )

    others = by("other")
    _naac_sheet(
        workbook, "Awards and recognition", f"Awards, fellowships and other recognition — {period}",
        ["Name of the teacher", "Kind", "Title of the award / activity", "Awarding agency / organisation",
         "Level", "Date"],
        [[
            teacher(r), d(r, "kind"), r.get("title"), d(r, "awarded_by"), d(r, "level"),
            _to_date(r.get("record_date")),
        ] for r in others],
    )
    return _save(workbook)


# ------------------------------------------------------------ import

def _headers(category: Category) -> list[tuple[str, str]]:
    """(header text, key) for a category's import sheet, in column order."""
    columns = [
        (f"{EMAIL_HEADER} *", "email"),
        (f"{category.title_label} *", "title"),
        (f"{category.date_label} * (YYYY-MM-DD)", "record_date"),
    ]
    if len(category.outcomes) > 1:
        columns.append(("Status", "outcome_status"))
    for item in category.fields:
        label = item.label + (" *" if item.required else "")
        if item.type == "multiselect":
            label += " (separate with ;)"
        elif item.type == "date":
            label += " (YYYY-MM-DD)"
        columns.append((label, f"details.{item.key}"))
    columns += [("Link", "link"), ("Abstract", "abstract")]
    return columns


def _header_key(text: str) -> str:
    return re.sub(r"\s*\((?:YYYY-MM-DD|separate with ;)\)\s*|\s*\*\s*", "", str(text or "")).strip().casefold()


def import_template() -> bytes:
    workbook = Workbook()
    guide = workbook.active
    guide.title = "Read me"
    lines = [
        "Research records import",
        "",
        "One sheet per kind of research. Fill in one row per record; leave a sheet empty if you have none of that kind.",
        f"'{EMAIL_HEADER}' is the email address of the teacher's Campus Capture account — it decides whose portfolio the record joins.",
        "Columns marked * are needed before a record can be sent to the Dean. Imported records arrive as drafts, so "
        "you can attach evidence and fill gaps before submitting.",
        "Dates are YYYY-MM-DD (or real Excel dates). Where a column has a drop-down, pick from it.",
        "Do not rename the sheets or the header row. Up to 500 rows per file.",
        "",
        "Choices for each drop-down column:",
    ]
    for line in lines:
        guide.append([line])
    guide["A1"].font = _TITLE_FONT
    for category in CATEGORIES.values():
        if len(category.outcomes) > 1:
            guide.append([f"{category.short} · Status: " + "; ".join(label for _, label in category.outcomes)])
        for item in category.fields:
            if item.options:
                guide.append([f"{category.short} · {item.label}: " + "; ".join(item.options)])
    guide.column_dimensions["A"].width = 130

    for category in CATEGORIES.values():
        sheet = workbook.create_sheet(_sheet_name(category.short))
        headers = _headers(category)
        sheet.append([text for text, _ in headers])
        for index, (_, key) in enumerate(headers, start=1):
            options: list[str] = []
            if key == "outcome_status":
                options = [label for _, label in category.outcomes]
            elif key.startswith("details."):
                item = next(f for f in category.fields if f.key == key[8:])
                if item.type == "select":
                    options = list(item.options)
            formula = '"' + ",".join(options) + '"'
            if options and len(formula) <= 255 and not any("," in o for o in options):
                column = get_column_letter(index)
                rule = DataValidation(type="list", formula1=formula, allow_blank=True)
                rule.add(f"{column}2:{column}{MAX_IMPORT_ROWS + 1}")
                sheet.add_data_validation(rule)
        _finish(sheet, 1)
        for index in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 24
    return _save(workbook)


class ImportFileError(ValueError):
    pass


def _plain(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        return value.strip() or None
    return value


def _date_text(value) -> str | None:
    """A date cell or text as YYYY-MM-DD; anything else is returned as-is
    for the router to report."""
    value = _plain(value)
    if value is None:
        return None
    text = str(value).strip()
    match = re.fullmatch(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})", text)
    if match:
        day, month, year = (int(g) for g in match.groups())
        try:
            return date(year, month, day).isoformat()
        except ValueError:
            return text
    return text


def parse_import(content: bytes) -> list[dict]:
    """Rows from a filled-in template: {"sheet", "row", "category", ...raw values}."""
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 -- openpyxl raises a variety of errors for non-xlsx input
        raise ImportFileError("That file could not be read as an Excel workbook (.xlsx). Use the template.") from None

    by_sheet = {category.short.casefold(): category for category in CATEGORIES.values()}
    by_sheet.update({category.label.casefold(): category for category in CATEGORIES.values()})
    rows: list[dict] = []
    for sheet in workbook.worksheets:
        category = by_sheet.get(sheet.title.strip().casefold())
        if category is None:
            continue
        known = {_header_key(text): key for text, key in _headers(category)}
        header: list[str | None] = []
        for number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            if number == 1:
                header = [known.get(_header_key(v)) if v is not None else None for v in values]
                if "title" not in header or "email" not in header:
                    raise ImportFileError(
                        f"The '{sheet.title}' sheet's header row has been changed. Download a fresh template."
                    )
                continue
            cells = {key: values[i] for i, key in enumerate(header) if key and i < len(values)}
            if all(_plain(v) in (None, "") for v in cells.values()):
                continue
            details: dict = {}
            for key, value in cells.items():
                if not key.startswith("details."):
                    continue
                field_key = key[8:]
                item = next(f for f in category.fields if f.key == field_key)
                if item.type == "multiselect":
                    text = _plain(value)
                    details[field_key] = [p.strip() for p in re.split(r"[;,]", str(text)) if p.strip()] if text else []
                elif item.type == "date":
                    details[field_key] = _date_text(value)
                else:
                    plain = _plain(value)
                    details[field_key] = str(plain) if isinstance(plain, (int, float)) and item.type != "number" else plain
            rows.append({
                "sheet": sheet.title,
                "row": number,
                "category": category.key,
                "email": str(_plain(cells.get("email")) or "").strip().lower(),
                "title": _plain(cells.get("title")),
                "record_date": _date_text(cells.get("record_date")),
                "outcome_status": _plain(cells.get("outcome_status")),
                "details": details,
                "link": _plain(cells.get("link")),
                "abstract": _plain(cells.get("abstract")),
            })
            if len(rows) > MAX_IMPORT_ROWS:
                raise ImportFileError(f"A file can hold at most {MAX_IMPORT_ROWS} records. Split it into smaller files.")
    workbook.close()
    if not rows:
        raise ImportFileError("No records found. Fill in at least one row on the category sheets of the template.")
    return rows
