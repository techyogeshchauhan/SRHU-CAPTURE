"""Excel (.xlsx) downloads of the lists people see: events, teachers.

One styled sheet per download: a title, a note of what it covers, a header
row and the rows. Every cell is written as data, never as a formula -- a
title such as =HYPERLINK(...) typed into a form must not run in whoever
opens the file (the CSV export guards against the same thing with csv_cell).
"""

from __future__ import annotations

import re
from datetime import date, datetime
from io import BytesIO
from typing import Any, Iterable

from fastapi import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.report_pdf import split_description
from app.services.storage_service import content_disposition

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_BRAND = "0D3668"
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor=_BRAND)
_TITLE_FONT = Font(bold=True, size=13, color=_BRAND)
_NOTE_FONT = Font(italic=True, color="5B6475")
_WRAP = Alignment(wrap_text=True, vertical="top")
_ISO_DAY = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def as_text_only(sheet) -> None:
    """Store anything openpyxl took for a formula as the text it is."""
    for row in sheet.iter_rows():
        for cell in row:
            if cell.data_type == "f":
                cell.data_type = "s"


def excel_value(value: Any) -> Any:
    """A value as a cell wants it: days as dates, lists joined, no timezones
    (Excel has none, and openpyxl refuses an aware datetime)."""
    if value is None or value == "":
        return None
    if isinstance(value, (list, tuple, set)):
        return ", ".join(str(item) for item in value if item not in (None, ""))
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        if _ISO_DAY.match(text):
            try:
                return date.fromisoformat(text)
            except ValueError:
                return text
        if len(text) > 10 and text[:4].isdigit() and text[4] == "-" and "T" in text:
            try:
                return datetime.fromisoformat(text.replace("Z", "+00:00")).replace(tzinfo=None)
            except ValueError:
                return text
        return text
    if isinstance(value, (int, float, bool)):
        return value
    return str(value)


def build_workbook(title: str, headers: list[str], rows: Iterable[list[Any]], *, note: str | None = None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = re.sub(r"[\[\]:*?/\\]", "-", title)[:31]
    sheet["A1"] = title
    sheet["A1"].font = _TITLE_FONT
    sheet["A2"] = note or f"Exported {datetime.now().strftime('%d %b %Y, %H:%M')}"
    sheet["A2"].font = _NOTE_FONT
    header_row = 4
    sheet.append([])
    sheet.append(headers)
    count = 0
    for row in rows:
        sheet.append([excel_value(value) for value in row])
        count += 1
    if count == 0:
        sheet.append(["Nothing matches these filters."])

    for cell in sheet[header_row]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    for index, column in enumerate(sheet.iter_cols(min_row=header_row), start=1):
        longest = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        sheet.column_dimensions[get_column_letter(index)].width = min(max(10, longest + 2), 60)
        for cell in column[1:]:
            if isinstance(cell.value, datetime):
                cell.number_format = "dd-mmm-yyyy hh:mm"
            elif isinstance(cell.value, date):
                cell.number_format = "dd-mmm-yyyy"
            elif isinstance(cell.value, str) and len(cell.value) > 40:
                cell.alignment = _WRAP
    as_text_only(sheet)

    out = BytesIO()
    workbook.save(out)
    return out.getvalue()


def xlsx_response(content: bytes, file_name: str) -> Response:
    return Response(
        content=content,
        media_type=XLSX,
        headers={"Content-Disposition": content_disposition("attachment", file_name), "Cache-Control": "no-store"},
    )


def stamped(stem: str) -> str:
    return f"{stem}-{date.today().isoformat()}.xlsx"


# ------------------------------------------------------------ events

EVENT_HEADERS = [
    "Event name", "Event type", "Status", "Start date", "End date", "Start time", "End time",
    "Venue", "Organiser", "Coordinator", "Coordinator contact", "Department",
    "Submitted by", "Submitter email", "Created on", "Description",
]


def _status_label(status: str | None) -> str:
    return (status or "").replace("_", " ").capitalize()


def event_rows(event_list: list[dict]) -> list[list[Any]]:
    """Rows for EVENT_HEADERS. The events have been through
    with_legacy_metadata, and attach_teachers where the submitter is wanted."""
    rows = []
    for event in event_list:
        description, meta = split_description(event.get("description"))
        rows.append([
            event.get("event_name"),
            event.get("event_type"),
            _status_label(event.get("status")),
            event.get("event_date"),
            event.get("end_date"),
            event.get("start_time"),
            event.get("end_time"),
            event.get("location"),
            event.get("organizer"),
            event.get("coordinator_name"),
            event.get("coordinator_contact"),
            event.get("department") or meta.get("department"),
            event.get("teacher_name") or event.get("owner_name"),
            event.get("teacher_email") or event.get("owner_email"),
            event.get("created_at"),
            description,
        ])
    return rows
